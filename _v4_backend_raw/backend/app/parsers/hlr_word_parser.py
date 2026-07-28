# -*- coding: utf-8 -*-
"""HLR Word parser — extracts software high-level requirements from .docx tables."""

from __future__ import annotations

from pathlib import Path

from docx import Document  # type: ignore

from app.models import HLRGlossaryEntry, HLROutput, HLRRequirement


# Expected labels for each row of requirement tables (for validation)
_REQ_FIELD_LABELS = [
    "需求ID",
    "需求中文",
    "对象类型",
    "是否衍生",
    "基本原理",
    "安全相关",
    "验证方法",
    "实现方法",
]


def _cell_text(table, row: int, col: int) -> str:
    """Extract stripped text from a table cell."""
    try:
        return table.cell(row, col).text.strip()
    except (IndexError, AttributeError):
        return ""


def _extract_requirement(table, source_name: str) -> HLRRequirement:
    """Extract one HLRRequirement from an 8×2 table."""
    cells = [_cell_text(table, r, 1) for r in range(8)]
    return HLRRequirement(
        requirement_id=cells[0],
        content=cells[1],
        object_type=cells[2],
        is_derived=cells[3],
        rationale=cells[4],
        is_safety_related=cells[5],
        verification_method=cells[6],
        implementation_method=cells[7],
        source_file=source_name,
    )


class HLRWordParser:
    """Parse a Software HLR Word document into structured requirements.

    Expects the first table to be a glossary (3 columns), followed by
    8-row × 2-column requirement tables.
    """

    def __init__(self, source_path: Path):
        self.source_path = source_path
        self.source_name = source_path.name

    def parse(self) -> HLROutput:
        """Main entry point. Parse the Word document and return HLROutput."""
        doc = Document(str(self.source_path))
        tables = doc.tables

        glossary: list[HLRGlossaryEntry] = []
        requirements: list[HLRRequirement] = []

        # Table 0 = glossary (3 cols: abbreviation, English full name, Chinese description)
        if len(tables) >= 1 and len(tables[0].columns) >= 3:
            glossary = self._parse_glossary(tables[0])

        # Tables 1..N = requirement tables (8 rows × 2 cols)
        for table in tables[1:]:
            if len(table.rows) >= 8 and len(table.columns) >= 2:
                req = _extract_requirement(table, self.source_name)
                # Skip tables with empty IDs (e.g., template rows)
                if req.requirement_id:
                    requirements.append(req)

        return HLROutput(
            source_file=self.source_name,
            total_count=len(requirements),
            requirements=requirements,
            glossary=glossary,
        )

    @staticmethod
    def _parse_glossary(table) -> list[HLRGlossaryEntry]:
        """Parse the glossary table (3 columns).

        First row is the header: 缩写, 英文全名, 中文说明
        """
        entries: list[HLRGlossaryEntry] = []
        for r in range(1, len(table.rows)):
            abbr = _cell_text(table, r, 0)
            eng = _cell_text(table, r, 1)
            chn = _cell_text(table, r, 2)
            if abbr:
                entries.append(
                    HLRGlossaryEntry(
                        abbreviation=abbr,
                        english_name=eng,
                        chinese_description=chn,
                    )
                )
        return entries
