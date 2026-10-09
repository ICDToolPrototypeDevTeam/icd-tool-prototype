# -*- coding: utf-8 -*-
"""
EoICD 定向定位工具 v2（内存优化版）
=====================================
两遍扫描策略：
1. 第一遍 read_only=True 只读表头，判断哪些 Sheet 需要过滤
2. 第二遍 read_only=True 流式读取数据，收集保留的行
3. 创建新 Excel 写入保留的数据

解决大文件内存问题：全程 read_only=True，不修改原文件。
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path
from typing import Any

import openpyxl
import yaml


def _detect_side_boundaries(row_cells: list[Any]) -> tuple[int | None, int | None]:
    pub_col = None
    sub_col = None
    for i, val in enumerate(row_cells, start=1):
        if val is not None:
            s = str(val).strip()
            if s == "Publisher" and pub_col is None:
                pub_col = i
            elif s == "Subscriber" and sub_col is None:
                sub_col = i
    return pub_col, sub_col


def _detect_layers(row_cells: list[Any], col_start: int, col_end: int) -> list[dict]:
    layers = []
    current = None
    for i in range(col_start, col_end):
        val = row_cells[i - 1] if i - 1 < len(row_cells) else None
        if val is not None:
            if current is not None:
                current["end_col"] = i
                layers.append(current)
            current = {"layer_type": str(val).strip(), "start_col": i, "end_col": col_end}
    if current is not None:
        current["end_col"] = col_end
        layers.append(current)
    return layers


def _read_attr_names(row_cells: list[Any], layers: list[dict]) -> None:
    for layer in layers:
        attr_names = []
        for c in range(layer["start_col"], layer["end_col"]):
            val = row_cells[c - 1] if c - 1 < len(row_cells) else None
            attr_names.append(str(val).strip() if val is not None else "")
        layer["attr_names"] = attr_names


def parse_filter_key(filter_key: str) -> tuple[str, str]:
    if "." in filter_key:
        parts = filter_key.split(".", 1)
        return parts[0].strip(), parts[1].strip()
    return "", filter_key.strip()


def find_attr_column(layers: list[dict], layer_name: str, attr_name: str) -> tuple[int, int] | None:
    for idx, layer in enumerate(layers):
        if layer_name and layer["layer_type"] != layer_name:
            continue
        for offset, name in enumerate(layer.get("attr_names", [])):
            if name == attr_name:
                return idx, offset
    return None


def check_row_match(row_cells: list[Any], layers: list[dict], conditions: dict) -> bool:
    """判断一行是否满足定位条件；当前实现为 OR 逻辑。"""
    for filter_key, allowed_values in conditions.items():
        layer_name, attr_name = parse_filter_key(filter_key)
        pos = find_attr_column(layers, layer_name, attr_name)
        if pos is None:
            continue
        layer_idx, attr_offset = pos
        layer = layers[layer_idx]
        col_idx = layer["start_col"] + attr_offset
        if col_idx > len(row_cells):
            continue
        cell_val = row_cells[col_idx - 1]
        cell_str = str(cell_val).strip() if cell_val is not None else ""
        for allowed in allowed_values:
            if cell_str == str(allowed).strip():
                return True
    return False


def locate_eoicd(input_path: Path, config: dict, side: str, output_path: Path) -> None:
    print(f"扫描: {input_path}")

    side_filters = config.get("filters", {}).get(side, {})
    sheets_to_filter = {}  # sheet_name -> (conditions, layers)
    all_sheets_info = {}   # sheet_name -> needs_filter: bool

    # ===== 第一遍：read_only=True 只读表头，判断哪些 Sheet 需要过滤 =====
    wb_scan = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    for sheet_name in wb_scan.sheetnames:
        conditions = side_filters.get(sheet_name)
        if not conditions:
            all_sheets_info[sheet_name] = False
            print(f"  {sheet_name}: 无定位条件，跳过")
            continue

        ws = wb_scan[sheet_name]
        header = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            header.append(row)
            if i >= 2:
                break

        if len(header) < 3:
            all_sheets_info[sheet_name] = False
            print(f"  {sheet_name}: 行数不足，跳过")
            continue

        row1, row2, row3 = header[0], header[1], header[2]
        pub_col, sub_col = _detect_side_boundaries(row1)
        max_col = max(len(row1), len(row2), len(row3))
        if side == "subscriber":
            # Subscriber 表要按「消费方」软件过滤：只扫描 Subscriber 段
            # （sub_col 起），否则 Software.Name 会错误命中 Publisher 段（信号源 LRU），
            # 导致散落设备（如 HF_AFTBFAN1）因"发布方在白名单内"被误保留。
            col_start = sub_col if sub_col else 1
            col_end = max_col + 1
        else:
            # Publisher 表只扫描 Publisher 段（1 .. sub_col）
            col_start = 1
            col_end = sub_col if sub_col else (max_col + 1)
        layers = _detect_layers(row2, col_start, col_end)
        _read_attr_names(row3, layers)

        missing = []
        for fk in conditions.keys():
            ln, an = parse_filter_key(fk)
            if find_attr_column(layers, ln, an) is None:
                missing.append(fk)

        if missing:
            all_sheets_info[sheet_name] = False
            print(f"  {sheet_name}: 属性 {missing} 不存在，跳过")
            continue

        sheets_to_filter[sheet_name] = (conditions, layers)
        all_sheets_info[sheet_name] = True
        print(f"  {sheet_name}: 需要过滤")

    wb_scan.close()

    if not sheets_to_filter:
        print("\n无需要过滤的 Sheet，直接复制文件")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(input_path, output_path)
        print(f"输出: {output_path}")
        return

    # ===== 第二遍：read_only=True 流式读取，收集保留的行 =====
    print(f"\n读取数据...")
    wb_data = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    sheets_retained_rows = {}  # sheet_name -> list of rows (包括表头)

    for sheet_name in wb_data.sheetnames:
        ws = wb_data[sheet_name]

        if not all_sheets_info.get(sheet_name, False):
            # 不需要过滤：读取所有行
            rows = list(ws.iter_rows(values_only=True))
            sheets_retained_rows[sheet_name] = rows
            print(f"  {sheet_name}: 保留全部 {len(rows)} 行")
            continue

        # 需要过滤：流式读取，只保留匹配的行
        conditions, layers = sheets_to_filter[sheet_name]
        retained = []
        header_count = 0
        data_count = 0
        filtered_count = 0

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 3:
                retained.append(row)
                header_count += 1
                continue
            data_count += 1
            if check_row_match(row, layers, conditions):
                retained.append(row)
            else:
                filtered_count += 1

        sheets_retained_rows[sheet_name] = retained
        print(f"  {sheet_name}: 过滤 {filtered_count} 行，保留 {len(retained) - 3} 行（原始 {data_count} 行）")

    wb_data.close()

    # ===== 第三遍：创建新 Workbook，写入保留的行 =====
    print(f"\n写入新文件...")
    wb_out = openpyxl.Workbook()
    # 删除默认创建的 sheet
    if "Sheet" in wb_out.sheetnames:
        wb_out.remove(wb_out["Sheet"])

    for sheet_name in wb_data.sheetnames:
        rows = sheets_retained_rows[sheet_name]
        ws_out = wb_out.create_sheet(title=sheet_name)
        for row_data in rows:
            ws_out.append(row_data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_path)
    print(f"\n输出: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="EoICD 定向定位工具（内存优化版）")
    parser.add_argument("-i", "--input", required=True)
    parser.add_argument("-c", "--config", required=True)
    parser.add_argument("-s", "--side", choices=["publisher", "subscriber"], required=True)
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)

    locate_eoicd(Path(args.input), config, args.side, Path(args.output))


if __name__ == "__main__":
    main()
