# -*- coding: utf-8 -*-
"""把正向多 judge 共识分析报告导出为 Excel 附件（与 Word 报告内容同源）。

复用 multi_judge_report / multi_judge_summary / multi_judge_aggregate 的取数
逻辑，确保 Excel 与 docx 的「判定分布 / 星级分布 / 处置建议 / 分析明细」
完全一致。Excel 含：一致性分析总结、汇总页（全量明细，原「分析明细」改名）、
以及按问题类型（不一致 / 未落实 / 部分落实 / 已落实 / 覆盖缺口）拆分的分工作表，
均可筛选、可排序。

用法（在 正向检查集成 目录下，脚本已收拢到 multi_judge协同/ 子目录）：
    python multi_judge协同/gen_excel_report.py
产物：output/EoICD至HLR正向一致性多模型共识分析报告_<YYYYMMDD_HHMM>.xlsx
      文件名时间戳与同次 run 的 Word 报告一致（取自 summary.json 的 run_timestamp），便于配对。
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from multi_judge_aggregate import _consensus_label
from multi_judge_report import (
    DETAIL_HEADERS,
    SUB_CLASS_DISPLAY,
    _build_paragraphs,
    _format_attrs,
    _star_str,
)
from multi_judge_summary import (
    AGREEMENT_CN,
    VERDICT_COLORS,
    VERDICT_GROUPS,
)

# 本脚本已收拢到 multi_judge协同/ 子目录，故 HERE 上溯一级才是工程根（runs/、output/ 所在）。
HERE = Path(__file__).resolve().parent.parent
# 数据类中间产物（聚合/汇总）下沉到 runs 工作区；xlsx 交付物落到 output/
AGG = HERE / "runs" / "integration_workspace" / "multi_judge_aggregate.json"
SUM = HERE / "runs" / "integration_workspace" / "multi_judge_summary.json"
# xlsx 文件名带运行时间戳，与同次 run 的 Word 报告配对（时间戳取自 summary.json 的 run_timestamp）
REPORT_STEM = "EoICD至HLR正向一致性多模型共识分析报告"

# 颜色（ARGB，与 Word 报告一致）
_VERDICT_FONT = {k: "FF" + v for k, v in VERDICT_COLORS.items()}
_GREEN = "FF008000"
_YELLOW = "FFCC8800"
_RED = "FFCC0000"
_GRAY = "FF808080"
_HEADER_FILL = PatternFill("solid", fgColor="FFD9E2F3")
_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True, color="FF1F3864")
_WRAP = Alignment(vertical="top", wrap_text=True)
_TOP = Alignment(vertical="top")


def _star_font(star: int) -> str:
    if star == 5:
        return _GREEN
    if star in (3, 4):
        return _YELLOW
    if star in (1, 2):
        return _RED
    return _GRAY


def _style_header_row(ws, ncols: int, row: int = 1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _set_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def sheet_summary_block(wb, summary: dict):
    """一、一致性分析总结：总览 + 判定分布 + 星级分布 + 处置建议。"""
    ws = wb.active
    ws.title = "一致性分析总结"
    r = 1
    ws.cell(r, 1, "EoICD 至 HLR 正向一致性多模型共识分析（Excel 版）").font = _TITLE_FONT
    r += 1
    ws.cell(r, 1, f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(italic=True, color="FF808080")
    r += 1
    ws.cell(r, 1, f"对 {summary.get('total', 0)} 条 EoICD 信号进行多 judge 并行独立判定，按共识度给出星级评价；"
                  f"判定粒度为「信号」（一个信号可命中多条 HLR 需求）。平均星级 {summary.get('average_star_rating', 0):.2f}。")
    r += 2

    # 判定分布
    ws.cell(r, 1, "判定分布").font = Font(bold=True, size=12)
    r += 1
    ws.cell(r, 1, "判定结果"); ws.cell(r, 2, "数量")
    _style_header_row(ws, 2, r)
    dist = summary.get("status_distribution", {})
    head = r
    r += 1
    for label, count in dist.items():
        ws.cell(r, 1, label).font = Font(bold=True, color=_VERDICT_FONT.get(label, _GRAY))
        ws.cell(r, 2, count)
        r += 1
    ws.cell(r, 1, "合计").font = Font(bold=True)
    ws.cell(r, 2, summary.get("total", 0)).font = Font(bold=True)
    r += 1
    # #12 未落实二级子类（对齐单 judge：已识别未承接 / 未识别）
    _sd = summary.get("status_distribution_detail") or {}
    if _sd and any(_sd.values()):
        ws.cell(r, 1, "  未落实子类").font = Font(bold=True, color=_VERDICT_FONT.get("未落实", _GRAY))
        r += 1
        for _k, _v in _sd.items():
            if _v:
                ws.cell(r, 1, "    %s" % SUB_CLASS_DISPLAY.get(_k, _k))
                ws.cell(r, 2, _v)
                r += 1
    r += 1

    # 星级分布
    ws.cell(r, 1, "星级分布").font = Font(bold=True, size=12)
    r += 1
    for i, h in enumerate(["星级", "共识等级", "数量", "说明"], start=1):
        ws.cell(r, i, h)
    _style_header_row(ws, 4, r)
    r += 1
    _star_desc = {
        "5": "全部 judge 判断完全一致，无字段级分歧",
        "4": "judge 判断一致，但存在字段异议",
        "3": "多数 judge 判断一致，少数意见仅涉及辅助字段",
        "2": "多数 judge 判断一致，少数意见涉及关键字段",
        "1": "judge 之间分歧或全部失效，需人工复核",
    }
    sd = summary.get("star_distribution", {})
    ad = summary.get("agreement_distribution", {})
    for sk in ("5", "4", "3", "2"):
        ws.cell(r, 1, _star_str(int(sk))).font = Font(bold=True, color=_YELLOW)
        ws.cell(r, 2, _consensus_label(int(sk), "")).font = Font(bold=True)
        ws.cell(r, 3, sd.get(sk, 0))
        ws.cell(r, 4, _star_desc[sk]).alignment = _WRAP
        r += 1
    for agr_key, label in (("split", "三方分歧"), ("single_source", "仅单源"), ("no_consensus", "全部失效")):
        ws.cell(r, 1, _star_str(1)).font = Font(bold=True, color=_YELLOW)
        ws.cell(r, 2, label).font = Font(bold=True)
        ws.cell(r, 3, ad.get(agr_key, 0))
        ws.cell(r, 4, AGREEMENT_CN.get(agr_key, label)).alignment = _WRAP
        r += 1
    ws.cell(r, 2, "平均星级").font = Font(bold=True)
    ws.cell(r, 3, round(summary.get("average_star_rating", 0), 2)).font = Font(bold=True)
    r += 2

    # 处置建议
    ws.cell(r, 1, "处置建议").font = Font(bold=True, size=12)
    r += 1
    ws.cell(r, 1, "建议"); ws.cell(r, 2, "处置")
    _style_header_row(ws, 2, r)
    r += 1
    for text, tag in _build_paragraphs():
        ws.cell(r, 1, text).alignment = _WRAP
        ws.cell(r, 2, tag).font = Font(bold=True)
        r += 1

    _set_widths(ws, {"A": 60, "B": 16, "C": 14, "D": 50})
    return ws


def _grouped(agg: dict) -> dict[str, list[tuple[str, dict]]]:
    """按 final_verdict 把信号分组（与 VERDICT_GROUPS 同序）。"""
    groups: dict[str, list[tuple[str, dict]]] = {g: [] for g in VERDICT_GROUPS}
    for full, item in agg.items():
        a = (item or {}).get("aggregation") or {}
        verdict = a.get("final_verdict") or "覆盖缺口"
        groups.setdefault(verdict, []).append((full, a))
    return groups


def sheet_detail(wb, agg: dict, summary: dict, title="汇总页", filter_verdict=None):
    """分析明细 / 按问题类型分类：一行 = 一个 EoICD 信号，可筛选/排序。

    title          工作表名（默认「汇总页」= 全量明细，原「分析明细」改名）。
    filter_verdict 仅当为 None 时输出全量；否则只输出该判定类型的信号。
    """
    ws = wb.create_sheet(title)
    for i, h in enumerate(DETAIL_HEADERS, start=1):
        ws.cell(1, i, h)
    _style_header_row(ws, len(DETAIL_HEADERS), 1)

    groups = _grouped(agg)
    if filter_verdict is not None:
        rows = groups.get(filter_verdict) or []
    else:
        rows = []
        for g in VERDICT_GROUPS:
            rows.extend(groups.get(g) or [])

    seq = 0
    r = 2
    for full, a in rows:
        seq += 1
        star = a.get("star_rating") or 0
        verdict = a.get("final_verdict") or filter_verdict or "覆盖缺口"
        ws.cell(r, 1, seq)
        ws.cell(r, 2, full).alignment = _WRAP
        c3 = ws.cell(r, 3, verdict)
        c3.font = Font(bold=True, color=_VERDICT_FONT.get(verdict, _GRAY))
        c3.alignment = _TOP
        # #12 子分类列（对齐单 judge：已识别未承接 / 未识别）
        ws.cell(r, 4, SUB_CLASS_DISPLAY.get(a.get("sub_classification")) or "—").alignment = _WRAP
        ws.cell(r, 5, a.get("matched_hlr_consensus") or "—").alignment = _WRAP
        ws.cell(r, 6, _format_attrs(a.get("inconsistent_attributes"))).alignment = _WRAP
        c6 = ws.cell(r, 7, a.get("consensus_label") or _consensus_label(star, ""))
        c6.font = Font(color=_star_font(star))
        c6.alignment = _TOP
        c7 = ws.cell(r, 8, _star_str(star))
        c7.font = Font(bold=True, color=_YELLOW)
        c7.alignment = _TOP
        ws.cell(r, 9, (a.get("analysis_summary") or "").strip() or "—").alignment = _WRAP
        r += 1

    last = r - 1
    ws.auto_filter.ref = f"A1:I{last}"
    ws.freeze_panes = "A2"
    _set_widths(ws, {"A": 6, "B": 42, "C": 10, "D": 34, "E": 34, "F": 22, "G": 18, "H": 8, "I": 50})
    return ws


def main():
    agg = json.loads(AGG.read_text(encoding="utf-8"))
    summary = json.loads(SUM.read_text(encoding="utf-8"))

    wb = Workbook()
    sheet_summary_block(wb, summary)
    # 汇总页：全量明细（原「分析明细」改名，内容不变）
    sheet_detail(wb, agg, summary, title="汇总页")
    # 按问题类型分类的分工作表：仅创建有数据的类型，列与汇总页一致
    grouped = _grouped(agg)
    for v in VERDICT_GROUPS:
        if grouped.get(v):
            sheet_detail(wb, agg, summary, title=v, filter_verdict=v)

    # 复用同次 run 的时间戳（runner 已写入 summary.json）；缺失则回退到当前时间
    run_ts = (summary or {}).get("run_timestamp") or datetime.now().strftime("%Y%m%d_%H%M")
    out_path = HERE / "output" / f"{REPORT_STEM}_{run_ts}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"✅ Excel 已写出: {out_path}")
    print(f"   工作表: {wb.sheetnames}")
    print(f"   明细信号数: {summary.get('total', 0)}")


if __name__ == "__main__":
    main()
